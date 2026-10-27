from django.contrib import admin
from .models import Pedido, DetallesPedido, BackOrder, Almacen
from productos.models import Producto

import datetime
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def _fecha(f):
    """Excel no acepta fechas con zona horaria: la pasamos a hora de México."""
    if isinstance(f, datetime.datetime) and timezone.is_aware(f):
        return timezone.localtime(f).replace(tzinfo=None)
    return f or ''


def generar_excel(nombre, encabezados, filas):
    wb = Workbook()
    ws = wb.active
    ws.title = nombre[:31]
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='1F3A5F')
    for fila in filas:
        ws.append(fila)
    for row in ws.iter_rows(min_row=2):
        for celda in row:
            if isinstance(celda.value, datetime.datetime):
                celda.number_format = 'DD/MM/YYYY HH:MM'
    for col in ws.columns:
        ancho = max(len(str(c.value or '')) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(ancho + 2, 60)
    ws.freeze_panes = 'A2'

    hoy = timezone.localtime().strftime('%Y-%m-%d')
    resp = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="{nombre}_{hoy}.xlsx"'
    wb.save(resp)
    return resp


@admin.action(description='📥 Exportar a Excel')
def exportar_pedidos(modeladmin, request, queryset):
    encabezados = ['Pedido', 'Proveedor', 'Estado', 'Fecha', 'Orden compra', 'Orden venta',
                   'Responsable', 'Producto', 'Descripción', 'Almacén', 'Cantidad',
                   'Recibido', 'Pendiente', 'Folio factura']
    filas = []
    pedidos = queryset.select_related('proveedor', 'responsable').prefetch_related(
        'detalles__producto', 'detalles__almacen')
    for p in pedidos:
        base = [p.numero_pedido, p.proveedor.nombre, p.get_estado_display(),
                _fecha(p.fecha_pedido), p.orden_compra or '', p.orden_venta or '',
                str(p.responsable or '')]
        detalles = list(p.detalles.all())
        if not detalles:
            filas.append(base + [''] * 7)
        for d in detalles:
            filas.append(base + [
                d.producto.nombre if d.producto else '',
                (d.producto.descripcion or '') if d.producto else '',
                str(d.almacen or ''), d.cantidad, d.recibido, d.pendiente,
                d.folio_factura or '',
            ])
    return generar_excel('Pedidos', encabezados, filas)


@admin.action(description='📥 Exportar a Excel')
def exportar_backorders(modeladmin, request, queryset):
    encabezados = ['Proveedor', 'Pedido', 'Orden compra', 'Producto', 'Descripción',
                   'Almacén', 'Cantidad', 'Entregado', 'Faltante', 'Resuelto', 'Fecha']
    filas = []
    for b in queryset.select_related('pedido__proveedor', 'producto', 'almacen'):
        filas.append([
            b.pedido.proveedor.nombre, b.pedido.numero_pedido, b.orden_compra or '',
            b.producto.nombre if b.producto else '',
            (b.producto.descripcion or '') if b.producto else '',
            str(b.almacen or ''), b.cantidad, b.entregado, b.cantidad - b.entregado,
            'Sí' if b.resuelto else 'No', _fecha(b.fecha_creacion),
        ])
    return generar_excel('BackOrders', encabezados, filas)

class DetallesPedidoInline(admin.TabularInline):
    model = DetallesPedido
    extra = 1
    fields = ('producto', 'almacen', 'cantidad', 'precio_unitario', 'recibido', 'pendiente', 'folio_factura')
    readonly_fields = ('precio_unitario','pendiente')
    raw_id_fields = ('producto',)

    def formfield_for_foreignkey(self, db_field, request, **kwargs):
        if db_field.name == 'producto':
            kwargs['queryset'] = Producto.objects.all()
        return super().formfield_for_foreignkey(db_field, request, **kwargs)
    
    def get_readonly_fields(self, request, obj=None):
        readonly = ['precio_unitario']
        if obj and obj.estado != 'entregado':
            readonly.extend(['recibido','pendiente', 'folio_factura', ])
        return readonly
    
    def has_delete_permission(self, request, obj=None):
        # Solo se pueden quitar productos en estado Pendiente
        if obj is None or obj.estado == 'pendiente':
            return super().has_delete_permission(request, obj)
        return False

@admin.register(Almacen)
class AlmacenAdmin(admin.ModelAdmin):
    list_display = ('nombre', 'ubicacion', 'fecha_creacion')
    search_fields = ('nombre',)

@admin.register(Pedido)
class PedidoAdmin(admin.ModelAdmin):
    list_display = ('numero_pedido', 'proveedor', 'estado', 'fecha_pedido')
    list_filter = ('estado', 'fecha_pedido')
    search_fields = ('numero_pedido', 'proveedor__nombre')
    inlines = [DetallesPedidoInline]
    actions = [exportar_pedidos]
    
    def get_readonly_fields(self, request, obj=None):
        readonly = []
        if obj:
            if obj.estado == 'pendiente':
                readonly = ['orden_compra', 'orden_venta', 'responsable', 'numero_pedido', 'proveedor']
            elif obj.estado == 'confirmado':
                readonly = ['responsable', 'numero_pedido', 'proveedor']
            elif obj.estado == 'solicitado':
                readonly = ['numero_pedido', 'proveedor', 'orden_compra', 'orden_venta']
            elif obj.estado == 'entregado':
                readonly = ['numero_pedido', 'proveedor', 'orden_compra', 'orden_venta', 'responsable', 'estado']
        return readonly

@admin.register(DetallesPedido)
class DetallesPedidoAdmin(admin.ModelAdmin):
    list_display = ('pedido', 'producto', 'almacen', 'cantidad', 'recibido', 'pendiente')
    list_filter = ('pedido__fecha_pedido',)
    search_fields = ('pedido__numero_pedido', 'producto__nombre')

from django.contrib import admin
from .models import BackOrder

@admin.register(BackOrder)
class BackOrderAdmin(admin.ModelAdmin):
    list_display = ('get_proveedor', 'orden_compra', 'producto', 'almacen', 'cantidad', 'entregado', 'resuelto', 'fecha_creacion')
    list_filter = ('resuelto', 'fecha_creacion', 'pedido__proveedor__nombre')
    search_fields = ('pedido__numero_pedido', 'producto__nombre', 'pedido__proveedor__nombre')
    readonly_fields = ('get_proveedor', 'orden_compra', 'producto', 'almacen', 'cantidad')
    actions = [exportar_backorders]
    
    def get_proveedor(self, obj):
        return obj.pedido.proveedor
    get_proveedor.short_description = 'Proveedor'
